#!/usr/bin/env python3
"""
📊 GERADOR DE ESCALAS EXCEL
===========================

Módulo responsável por gerar planilhas Excel com escalas de funcionários
formatadas, incluindo intervalos de tempo e atividades alocadas.

Dependências:
    - openpyxl: pip install openpyxl
"""

import os
import json
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Set
from dataclasses import dataclass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_DISPONIVEL = True
except ImportError:
    OPENPYXL_DISPONIVEL = False

from services.exportacao.escalas.parser_logs_funcionarios import (
    ParserLogsFuncionarios,
    AtividadeAlocada
)


@dataclass
class ConfiguracaoExcel:
    """Configurações para geração da planilha Excel."""
    intervalo_minutos: int = 30
    cor_cabecalho: str = "4472C4"
    cor_atividade: str = "92D050"
    largura_coluna_intervalo: int = 22
    largura_coluna_funcionario: int = 18
    altura_linha: int = 35
    tamanho_fonte_atividade: int = 8
    exibir_nomes_completos: bool = True
    excluir_linhas_vazias: bool = True


class GeradorEscalaExcel:
    """
    Gerador de planilhas Excel com escalas de funcionários.

    Gera planilhas formatadas com:
    - Intervalos de tempo nas linhas
    - Funcionários nas colunas
    - Atividades alocadas nas células
    """

    def __init__(
        self,
        arquivo_funcionarios: str = "data/funcionarios/funcionarios.json",
        diretorio_logs: str = "logs/funcionarios/sucesso",
        diretorio_saida: str = "data/escalas"
    ):
        """
        Inicializa o gerador.

        Args:
            arquivo_funcionarios: Caminho para o JSON de funcionários
            diretorio_logs: Caminho para logs de alocação
            diretorio_saida: Diretório para salvar a planilha
        """
        if not OPENPYXL_DISPONIVEL:
            raise ImportError(
                "Módulo openpyxl não encontrado. "
                "Instale com: pip install openpyxl"
            )

        self.arquivo_funcionarios = arquivo_funcionarios
        self.diretorio_logs = diretorio_logs
        self.diretorio_saida = diretorio_saida
        self.config = ConfiguracaoExcel()

        self._funcionarios: Dict[int, dict] = {}
        self._parser = ParserLogsFuncionarios(diretorio_logs)

    def _carregar_funcionarios(self) -> bool:
        """
        Carrega funcionários do arquivo JSON.

        Returns:
            True se carregou com sucesso
        """
        if not os.path.exists(self.arquivo_funcionarios):
            return False

        with open(self.arquivo_funcionarios, 'r', encoding='utf-8') as f:
            data = json.load(f)

        self._funcionarios = {
            f['id']: f for f in data.get('funcionarios', [])
        }

        return len(self._funcionarios) > 0

    def _criar_estilos(self) -> dict:
        """
        Cria estilos para a planilha.

        Returns:
            Dicionário com estilos
        """
        return {
            'header_fill': PatternFill(
                start_color=self.config.cor_cabecalho,
                end_color=self.config.cor_cabecalho,
                fill_type="solid"
            ),
            'header_font': Font(bold=True, color="FFFFFF"),
            'atividade_fill': PatternFill(
                start_color=self.config.cor_atividade,
                end_color=self.config.cor_atividade,
                fill_type="solid"
            ),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }

    def _gerar_intervalos(
        self,
        inicio: datetime,
        fim: datetime
    ) -> List[datetime]:
        """
        Gera lista de intervalos de tempo.

        Args:
            inicio: Data/hora inicial
            fim: Data/hora final

        Returns:
            Lista de datetimes representando início de cada intervalo
        """
        intervalos = []
        atual = inicio.replace(
            minute=(inicio.minute // self.config.intervalo_minutos)
                   * self.config.intervalo_minutos,
            second=0,
            microsecond=0
        )

        while atual <= fim:
            intervalos.append(atual)
            atual += timedelta(minutes=self.config.intervalo_minutos)

        return intervalos

    def _formatar_intervalo(self, inicio: datetime) -> str:
        """
        Formata um intervalo para exibição.

        Args:
            inicio: Datetime do início do intervalo

        Returns:
            String formatada "HH:MM – HH:MM [DD/MM]"
        """
        fim = inicio + timedelta(minutes=self.config.intervalo_minutos)
        return f"{inicio.strftime('%H:%M')} – {fim.strftime('%H:%M')} [{inicio.strftime('%d/%m')}]"

    def _encontrar_intervalos_com_atividades(
        self,
        intervalos: List[datetime],
        atividades: Dict[str, List[AtividadeAlocada]]
    ) -> Set[int]:
        """
        Encontra índices de intervalos que possuem atividades.

        Args:
            intervalos: Lista de intervalos
            atividades: Dicionário de atividades por funcionário

        Returns:
            Conjunto de índices de intervalos com atividades
        """
        indices = set()

        for lista_ativ in atividades.values():
            for ativ in lista_ativ:
                for idx, intervalo in enumerate(intervalos):
                    # Verificar sobreposição entre atividade e intervalo
                    fim_intervalo = intervalo + timedelta(minutes=self.config.intervalo_minutos)
                    if ativ.inicio < fim_intervalo and ativ.fim > intervalo:
                        indices.add(idx)

        return indices

    def gerar_escala(
        self,
        nome_arquivo: str = "escala_funcionarios_atividades.xlsx"
    ) -> Optional[str]:
        """
        Gera a planilha Excel com a escala de funcionários.

        Args:
            nome_arquivo: Nome do arquivo de saída

        Returns:
            Caminho do arquivo gerado ou None se falhar
        """
        # Carregar dados
        if not self._carregar_funcionarios():
            print("❌ Erro ao carregar funcionários")
            return None

        if not self._parser.carregar_logs():
            print("❌ Erro ao carregar logs de funcionários")
            return None

        atividades = self._parser.obter_atividades_por_funcionario()
        inicio, fim = self._parser.obter_intervalo_temporal()

        if not inicio or not fim:
            print("❌ Nenhuma atividade encontrada nos logs")
            return None

        # Gerar intervalos
        todos_intervalos = self._gerar_intervalos(inicio, fim)

        # Filtrar intervalos vazios se configurado
        if self.config.excluir_linhas_vazias:
            indices_ativos = self._encontrar_intervalos_com_atividades(
                todos_intervalos, atividades
            )
            intervalos = [
                (idx, todos_intervalos[idx])
                for idx in sorted(indices_ativos)
            ]
        else:
            intervalos = list(enumerate(todos_intervalos))

        # Criar workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Escala de Funcionários"

        estilos = self._criar_estilos()

        # Cabeçalho - coluna de intervalo
        ws['A1'] = 'Intervalo'
        ws['A1'].fill = estilos['header_fill']
        ws['A1'].font = estilos['header_font']
        ws['A1'].border = estilos['border']

        # Cabeçalho - colunas de funcionários
        col = 2
        func_cols = {}

        for func_id in sorted(self._funcionarios.keys()):
            func = self._funcionarios[func_id]
            cell = ws.cell(row=1, column=col, value=func['nome'])
            cell.fill = estilos['header_fill']
            cell.font = estilos['header_font']
            cell.border = estilos['border']
            cell.alignment = Alignment(textRotation=90, horizontal='center')
            func_cols[func['nome']] = col
            col += 1

        # Preencher intervalos
        for row_idx, (orig_idx, intervalo) in enumerate(intervalos, start=2):
            cell = ws.cell(
                row=row_idx,
                column=1,
                value=self._formatar_intervalo(intervalo)
            )
            cell.border = estilos['border']
            cell.alignment = Alignment(horizontal='left')

        # Preencher atividades
        for nome_func, lista_ativ in atividades.items():
            if nome_func not in func_cols:
                continue

            col = func_cols[nome_func]

            for ativ in lista_ativ:
                for row_idx, (orig_idx, intervalo) in enumerate(intervalos, start=2):
                    # Verificar se há sobreposição entre atividade e intervalo
                    fim_intervalo = intervalo + timedelta(minutes=self.config.intervalo_minutos)
                    if ativ.inicio < fim_intervalo and ativ.fim > intervalo:
                        cell = ws.cell(row=row_idx, column=col)

                        # Nome da atividade
                        nome = ativ.nome_atividade
                        if not self.config.exibir_nomes_completos and len(nome) > 20:
                            nome = nome[:18] + '..'

                        cell.value = nome
                        cell.fill = estilos['atividade_fill']
                        cell.border = estilos['border']
                        cell.alignment = Alignment(
                            horizontal='center',
                            vertical='center',
                            wrap_text=True
                        )
                        cell.font = Font(size=self.config.tamanho_fonte_atividade)

        # Ajustar larguras
        ws.column_dimensions['A'].width = self.config.largura_coluna_intervalo

        for col in range(2, len(self._funcionarios) + 2):
            ws.column_dimensions[get_column_letter(col)].width = \
                self.config.largura_coluna_funcionario

        # Ajustar alturas
        for row in range(2, len(intervalos) + 2):
            ws.row_dimensions[row].height = self.config.altura_linha

        # Criar diretório se não existir
        os.makedirs(self.diretorio_saida, exist_ok=True)

        # Salvar
        caminho_saida = os.path.join(self.diretorio_saida, nome_arquivo)
        wb.save(caminho_saida)

        print(f"✅ Planilha salva em: {caminho_saida}")
        print(f"   Total de funcionários: {len(self._funcionarios)}")
        print(f"   Funcionários com atividades: {len(atividades)}")
        print(f"   Intervalos com atividades: {len(intervalos)}")

        return caminho_saida

    def configurar(self, **kwargs) -> 'GeradorEscalaExcel':
        """
        Configura opções do gerador.

        Args:
            **kwargs: Opções de configuração

        Returns:
            self para encadeamento
        """
        for key, value in kwargs.items():
            if hasattr(self.config, key):
                setattr(self.config, key, value)

        return self


def gerar_escala_funcionarios(
    arquivo_funcionarios: str = "data/funcionarios/funcionarios.json",
    diretorio_logs: str = "logs/funcionarios/sucesso",
    diretorio_saida: str = "data/escalas",
    nome_arquivo: str = "escala_funcionarios_atividades.xlsx",
    excluir_linhas_vazias: bool = True,
    nomes_completos: bool = True
) -> Optional[str]:
    """
    Função utilitária para gerar escala de funcionários.

    Args:
        arquivo_funcionarios: Caminho para JSON de funcionários
        diretorio_logs: Caminho para logs de alocação
        diretorio_saida: Diretório de saída
        nome_arquivo: Nome do arquivo Excel
        excluir_linhas_vazias: Se True, remove linhas sem atividades
        nomes_completos: Se True, exibe nomes completos das atividades

    Returns:
        Caminho do arquivo gerado ou None
    """
    gerador = GeradorEscalaExcel(
        arquivo_funcionarios=arquivo_funcionarios,
        diretorio_logs=diretorio_logs,
        diretorio_saida=diretorio_saida
    )

    gerador.configurar(
        excluir_linhas_vazias=excluir_linhas_vazias,
        exibir_nomes_completos=nomes_completos
    )

    return gerador.gerar_escala(nome_arquivo)


# Execução direta para testes
if __name__ == "__main__":
    caminho = gerar_escala_funcionarios()
    if caminho:
        print(f"\n📊 Escala gerada com sucesso: {caminho}")
    else:
        print("\n❌ Falha ao gerar escala")
